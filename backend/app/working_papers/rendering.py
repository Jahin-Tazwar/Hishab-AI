"""Render a working paper to .docx and .pdf.

Today supports `at_risk_itc_schedule`. Each recipe-kind has a dedicated
docx builder. PDF rendering reuses `app.notices.rendering.render_pdf`
which already handles LibreOffice headless conversion.

The .docx is a CA working paper: firm letterhead, a unique reference, the
at-risk comparison (claimed vs supplier-reported VAT + variance), footing
totals, a prepared-by / reviewed-by sign-off block, and page numbers.
"""
from __future__ import annotations

import io
import re
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.notices.rendering import PdfRendererUnavailableError, render_pdf


# ── helpers ──────────────────────────────────────────────────────────────


def _fmt_bdt(v: Any) -> str:
    """Format an amount with South-Asian (lakh/crore) digit grouping + 2dp.

    Mirrors the frontend `formatBDT` so the on-screen and exported numbers
    match. Returns "" for None/empty so blank cells stay blank.
    """
    if v is None or v == "":
        return ""
    d = Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    neg = d < 0
    d = abs(d)
    int_part, dec_part = f"{d:.2f}".split(".")
    if len(int_part) <= 3:
        grouped = int_part
    else:
        last3 = int_part[-3:]
        rest = int_part[:-3]
        rest = re.sub(r"(?<=\d)(?=(?:\d\d)+$)", ",", rest)
        grouped = f"{rest},{last3}"
    return f"{'-' if neg else ''}{grouped}.{dec_part}"


def _right(cell) -> None:
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _bold_cell(cell) -> None:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True


def _page_number_footer(doc: Document) -> None:
    """Add a centered 'Page X of Y' field to the section footer."""
    footer = doc.sections[0].footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run("Page ")
    _field(para, "PAGE")
    para.add_run(" of ")
    _field(para, "NUMPAGES")


def _field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr_el)
    run._r.append(end)


def _notes_blocks(notes_html: str) -> list[str]:
    """Extract block-level text from sanitized notes HTML, preserving the
    paragraph/list structure (TipTap emits <p>/<li> with no newlines, so a
    flat text_content() would collapse everything into one run-on block)."""
    from lxml import html as lxml_html

    try:
        root = lxml_html.fromstring(f"<div>{notes_html}</div>")
    except Exception:
        return [notes_html.strip()] if notes_html.strip() else []

    blocks: list[str] = []
    for el in root.iter("p", "li"):
        txt = el.text_content().strip()
        if txt:
            prefix = "•  " if el.tag == "li" else ""
            blocks.append(prefix + txt)
    if not blocks:  # inline-only HTML (e.g. just "<b>x</b>")
        txt = root.text_content().strip()
        if txt:
            blocks.append(txt)
    return blocks


# ── public API ───────────────────────────────────────────────────────────


def render_docx(
    *,
    payload: dict[str, Any],
    notes_html: str,
    tenant: dict[str, Any],
    meta: Optional[dict[str, Any]] = None,
) -> bytes:
    kind = payload.get("kind")
    if kind == "at_risk_itc_schedule":
        return _render_at_risk_itc_schedule(payload, notes_html, tenant, meta or {})
    raise ValueError(f"Unknown working paper kind: {kind!r}")


def _emit_summary_table(doc, s: dict) -> None:
    summary_table = doc.add_table(rows=4, cols=2)
    summary_table.style = "Light List"
    pairs = [
        ("Total VAT claimed (BDT)", _fmt_bdt(s["total_vat_claimed_bdt"])),
        ("Safe ITC (BDT)", _fmt_bdt(s["safe_itc_bdt"])),
        ("At-risk ITC (BDT)", _fmt_bdt(s["at_risk_itc_bdt"])),
        ("Lines flagged",
         f"{s['at_risk_line_count']} of {s['total_lines']} "
         f"(across {s['supplier_count_at_risk']} supplier(s))"),
    ]
    for i, (k, v) in enumerate(pairs):
        summary_table.rows[i].cells[0].text = k
        summary_table.rows[i].cells[1].text = v
        _right(summary_table.rows[i].cells[1])


def _emit_supplier_section(doc, supplier_groups: list) -> Decimal:
    grand_total = Decimal("0.00")
    for grp in supplier_groups:
        grand_total += Decimal(str(grp.get("total_vat_at_risk_bdt") or "0"))
        gh = doc.add_paragraph()
        gr = gh.add_run(
            f"{grp.get('supplier_name') or '(unknown supplier)'}"
            + (f" — BIN {grp['supplier_bin']}" if grp.get("supplier_bin") else "")
            + f"   ·   VAT at risk: BDT {_fmt_bdt(grp['total_vat_at_risk_bdt'])}"
            + f"   ·   {grp['line_count']} line(s)")
        gr.bold = True

        lines = grp["lines"]
        table = doc.add_table(rows=2 + len(lines), cols=7)
        table.style = "Light List"
        headers = ["Invoice no", "Invoice date", "Claimed VAT (BDT)",
                   "Supplier VAT (BDT)", "Variance (BDT)", "Match",
                   "Recommended action"]
        for c, label in enumerate(headers):
            table.rows[0].cells[c].text = label
        for c in (2, 3, 4):
            _right(table.rows[0].cells[c])
        for i, line in enumerate(lines, start=1):
            row = table.rows[i].cells
            row[0].text = str(line.get("invoice_no") or "")
            row[1].text = str(line.get("invoice_date") or "")
            row[2].text = _fmt_bdt(line.get("vat_amount_bdt"))
            sf_vat = line.get("sf_vat_amount_bdt")
            row[3].text = _fmt_bdt(sf_vat) if sf_vat else "not filed"
            row[4].text = _fmt_bdt(line.get("vat_variance_bdt"))
            override = line.get("ca_override")
            ms = line["match_status"]
            match_txt = ms if not override else f"{ms} ({override})"
            reason = line.get("discrepancy_reason")
            row[5].text = f"{match_txt}\n{reason}" if reason else match_txt
            row[6].text = (line["recommended_action"] or "").replace("_", " ")
            for c in (2, 3, 4):
                _right(row[c])
        foot = table.rows[1 + len(lines)].cells
        foot[1].text = "Subtotal"
        _bold_cell(foot[1])
        foot[4].text = _fmt_bdt(grp["total_vat_at_risk_bdt"])
        _right(foot[4]); _bold_cell(foot[4])
    return grand_total


def _render_at_risk_itc_schedule(
    payload: dict, notes_html: str, tenant: dict, meta: dict,
) -> bytes:
    doc = Document()

    # ── Letterhead ──
    firm = tenant.get("firm_name") or "Chartered Accountants"
    p = doc.add_paragraph()
    r = p.add_run(firm)
    r.bold = True
    r.font.size = Pt(14)
    if tenant.get("firm_name_bn"):
        bn = doc.add_paragraph(tenant["firm_name_bn"])
        bn.runs[0].font.size = Pt(11)
    contact_bits = [
        tenant.get("address"),
        f"ICAB Reg: {tenant['icab_reg_no']}" if tenant.get("icab_reg_no") else None,
        tenant.get("email"),
        tenant.get("phone"),
    ]
    contact = "  ·  ".join(b for b in contact_bits if b)
    if contact:
        cp = doc.add_paragraph(contact)
        cp.runs[0].font.size = Pt(9)

    # ── Reference / generated / status block ──
    ref_line = []
    if meta.get("reference"):
        ref_line.append(f"Ref: {meta['reference']}")
    if meta.get("generated_on"):
        ref_line.append(f"Generated: {meta['generated_on']}")
    if meta.get("status"):
        ref_line.append(f"Status: {str(meta['status']).upper()}")
    if ref_line:
        rp = doc.add_paragraph("   ·   ".join(ref_line))
        rp.runs[0].font.size = Pt(9)
        rp.runs[0].italic = True

    # ── Title ──
    doc.add_paragraph("")
    t = doc.add_paragraph()
    tr = t.add_run("At-Risk Input Tax Credit Schedule")
    tr.bold = True
    tr.font.size = Pt(16)

    sh = doc.add_paragraph()
    sh.add_run(
        f"Client: {payload['client_name']}"
        + (f" (BIN {payload['client_bin']})" if payload.get("client_bin") else "")
    )
    sh2 = doc.add_paragraph()
    sh2.add_run(f"Period: {payload['period_start']} to {payload['period_end']}")

    # ── Summary ──
    doc.add_paragraph("")
    doc.add_paragraph().add_run("Summary").bold = True
    _emit_summary_table(doc, payload["summary"])

    # ── Per-supplier detail ──
    doc.add_paragraph("")
    doc.add_paragraph().add_run("At-risk lines by supplier").bold = True

    grand_total = _emit_supplier_section(doc, payload["supplier_groups"])

    # grand total
    gt = doc.add_paragraph()
    gtr = gt.add_run(f"Total at-risk ITC: BDT {_fmt_bdt(grand_total)}")
    gtr.bold = True
    gtr.font.size = Pt(12)

    # ── CA commentary (structure-preserving) ──
    blocks = _notes_blocks(notes_html) if notes_html else []
    if blocks:
        doc.add_paragraph("")
        doc.add_paragraph().add_run("CA commentary").bold = True
        for b in blocks:
            doc.add_paragraph(b)

    # ── Sign-off ──
    doc.add_paragraph("")
    signoff = doc.add_table(rows=2, cols=2)
    signoff.style = "Light List"
    signoff.rows[0].cells[0].text = "Prepared by"
    signoff.rows[0].cells[1].text = "Reviewed by"
    _bold_cell(signoff.rows[0].cells[0])
    _bold_cell(signoff.rows[0].cells[1])
    signoff.rows[1].cells[0].text = meta.get("prepared_by") or "— pending —"
    signoff.rows[1].cells[1].text = meta.get("reviewed_by") or "— pending —"

    _page_number_footer(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


__all__ = [
    "render_docx", "render_pdf", "_fmt_bdt", "PdfRendererUnavailableError",
    "_emit_summary_table", "_emit_supplier_section",
]
