"""Render a saved draft to .docx (always) and .pdf (when LibreOffice is on PATH).

The HTML body is parsed as a stream of <p>/<sup> tags using lxml. Citation
markers (<sup class="citation" data-ref="CIT-N">) are rendered inline as
superscript numerals; the canonical citation list is appended as a
separate section at the end of the document.

The .docx → .pdf step shells out to `soffice --headless --convert-to pdf`.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from typing import Any

from docx import Document
from docx.shared import Pt
from lxml import html as lxml_html


class PdfRendererUnavailableError(RuntimeError):
    """Raised when LibreOffice is not available on this host."""


def _soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("soffice.exe")


def _add_paragraph_from_html(doc, html_fragment: str) -> None:
    """Append one Word paragraph from a single <p>…</p> HTML fragment.

    Walks the fragment children:
    - text nodes → run
    - <sup class="citation"> → superscript run with its inner text
    - <strong>, <em> → bold / italic runs
    Anything else is rendered as plain text.
    """
    p = doc.add_paragraph()
    root = lxml_html.fromstring(html_fragment if html_fragment.strip().startswith("<")
                                else f"<p>{html_fragment}</p>")

    def _emit_text(text: str, *, bold=False, italic=False, superscript=False) -> None:
        if not text:
            return
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.superscript = superscript
        run.font.name = "Noto Sans Bengali"
        run.font.size = Pt(11)

    def _walk(node, *, bold=False, italic=False, superscript=False):
        if node.text:
            _emit_text(node.text, bold=bold, italic=italic, superscript=superscript)
        for child in node:
            tag = child.tag.lower() if isinstance(child.tag, str) else ""
            cb, ci, cs = bold, italic, superscript
            if tag == "strong":
                cb = True
            elif tag == "em":
                ci = True
            elif tag == "sup":
                cs = True
            _walk(child, bold=cb, italic=ci, superscript=cs)
            if child.tail:
                _emit_text(child.tail, bold=bold, italic=italic, superscript=superscript)

    _walk(root)


def _add_letterhead(doc, tenant: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    run = p.add_run(tenant.get("firm_name", "Chartered Accountants"))
    run.bold = True
    run.font.size = Pt(14)
    if tenant.get("address"):
        addr = doc.add_paragraph(tenant["address"])
        addr.runs[0].font.size = Pt(10)
    doc.add_paragraph("")


def _add_reference_block(doc, notice: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    p.add_run("প্রসঙ্গ: ").bold = True
    p.add_run(
        f"NBR Notice No. {notice.get('notice_no','(unknown)')} "
        f"dated {notice.get('notice_date','(unknown)')} — "
        f"VAT period {notice.get('period_start','?')} to {notice.get('period_end','?')}, "
        f"BIN {notice.get('taxpayer_bin','?')}."
    )
    doc.add_paragraph("")
    doc.add_paragraph("প্রিয় মহোদয়,")


def _add_appendix(doc, appendix_json: dict[str, Any]) -> None:
    doc.add_paragraph("")
    h = doc.add_paragraph()
    h.add_run("Appendix — Computation Summary (BDT)").bold = True

    rows = appendix_json.get("rows") or []
    if not rows:
        doc.add_paragraph("(no rows)")
        return
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Light List"
    hdr = table.rows[0].cells
    hdr[0].text = "Item"
    hdr[1].text = "Amount (BDT)"
    hdr[2].text = "Note"
    for i, r in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text = str(r.get("label", ""))
        cells[1].text = str(r.get("value_bdt", "") or "")
        cells[2].text = str(r.get("note", "") or "")


def _add_citations(doc, citations: list[dict[str, Any]]) -> None:
    if not citations:
        return
    doc.add_paragraph("")
    h = doc.add_paragraph()
    h.add_run("Citations").bold = True
    for i, c in enumerate(citations, start=1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(c.get("source_ref", "") + " — ").bold = True
        p.add_run(c.get("snippet", ""))


def render_docx(
    *, draft: dict[str, Any], notice: dict[str, Any], tenant: dict[str, Any],
) -> bytes:
    """Build a .docx reply and return its bytes."""
    doc = Document()
    _add_letterhead(doc, tenant)
    _add_reference_block(doc, notice)

    body_html = draft.get("body_html", "")
    if body_html:
        root = lxml_html.fromstring(f"<div>{body_html}</div>")
        for child in root.iter():
            if isinstance(child.tag, str) and child.tag.lower() == "p":
                _add_paragraph_from_html(
                    doc, lxml_html.tostring(child, encoding="unicode"),
                )

    doc.add_paragraph("")
    doc.add_paragraph("ধন্যবাদান্তে,")
    doc.add_paragraph(tenant.get("firm_name", ""))

    _add_appendix(doc, draft.get("appendix_json", {}))
    _add_citations(doc, draft.get("citations", []))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(docx_bytes: bytes) -> bytes:
    """Convert .docx → .pdf via LibreOffice headless.

    Raises PdfRendererUnavailableError if soffice is not on PATH.
    """
    soffice = _soffice_path()
    if not soffice:
        raise PdfRendererUnavailableError(
            "LibreOffice (soffice) not on PATH — PDF rendering unavailable"
        )
    with tempfile.TemporaryDirectory() as td:
        in_path = os.path.join(td, "in.docx")
        with open(in_path, "wb") as f:
            f.write(docx_bytes)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", td, in_path],
            check=True, capture_output=True, timeout=60,
        )
        out_path = os.path.join(td, "in.pdf")
        with open(out_path, "rb") as f:
            return f.read()
