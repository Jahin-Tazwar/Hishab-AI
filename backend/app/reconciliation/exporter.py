"""XLSX export of a completed reconciliation.

Produces a workbook with two sheets:
  1. Summary — headline aggregates (counts + ITC numbers)
  2. Line Items — per-invoice rows colour-coded by match_status

Pure function `build_xlsx_workbook` accepts in-memory dicts so it can be
unit-tested without touching the database. The async `export_reconciliation`
wrapper loads the rows via the service-role supabase client and serialises
the workbook to bytes.
"""
from __future__ import annotations

import asyncio
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.exceptions import (
    DocumentTenantMismatchError,
    ReconciliationNotFoundError,
)
from app.database import get_supabase_admin


# ── Status → fill colour (light, print-friendly) ────────────────────────
_STATUS_FILLS: dict[str, PatternFill] = {
    "exact":    PatternFill("solid", fgColor="C6EFCE"),  # green
    "fuzzy":    PatternFill("solid", fgColor="FFEB9C"),  # amber
    "partial":  PatternFill("solid", fgColor="FFD8B1"),  # orange
    "no_match": PatternFill("solid", fgColor="FFC7CE"),  # red
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

_LINE_HEADERS = [
    "PR Invoice No", "PR Supplier BIN", "PR Supplier Name", "PR Invoice Date",
    "PR Taxable (BDT)", "PR VAT (BDT)",
    "SF Invoice No", "SF Invoice Date", "SF Taxable (BDT)", "SF VAT (BDT)",
    "Match Status", "Match Score", "CA Override", "CA Notes",
]


def _style_header_row(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    ws.freeze_panes = "A2"


def _autosize(ws, *, max_width: int = 28) -> None:
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        widest = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            if length > widest:
                widest = length
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            widest + 2, max_width
        )


def build_xlsx_workbook(
    *,
    header: dict,
    line_items: list[dict],
) -> bytes:
    """Build an XLSX workbook in memory, return bytes.

    `header` is the vat_reconciliations row (dict). `line_items` is the
    list of recon_line_items rows (dicts) ordered by status severity then
    invoice date for readability.
    """
    wb = Workbook()

    # ── Summary sheet ─────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"

    rows = [
        ("Reconciliation ID", str(header.get("id", ""))),
        ("Period Start", str(header.get("period_start", ""))),
        ("Period End", str(header.get("period_end", ""))),
        ("Status", header.get("status", "")),
        ("", ""),
        ("Total Invoices", header.get("total_invoices", 0)),
        ("Matched Exact", header.get("matched_exact", 0)),
        ("Matched Fuzzy", header.get("matched_fuzzy", 0)),
        ("Partial Match", header.get("partial_match", 0)),
        ("No Match", header.get("no_match", 0)),
        ("", ""),
        ("Total VAT Claimed (BDT)", header.get("total_vat_claimed_bdt", 0)),
        ("Safe ITC (BDT)", header.get("safe_itc_bdt", 0)),
        ("At-Risk ITC (BDT)", header.get("at_risk_itc_bdt", 0)),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=label).font = Font(bold=bool(label))
        ws.cell(row=r, column=2, value=value)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 36

    # ── Line Items sheet ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Line Items")
    ws2.append(_LINE_HEADERS)
    _style_header_row(ws2, len(_LINE_HEADERS))

    # Sort: no_match first (most attention), then partial, fuzzy, exact
    severity = {"no_match": 0, "partial": 1, "fuzzy": 2, "exact": 3}
    sorted_items = sorted(
        line_items,
        key=lambda li: (
            severity.get(li.get("match_status", "no_match"), 4),
            str(li.get("pr_invoice_date") or ""),
        ),
    )

    for li in sorted_items:
        ws2.append([
            li.get("pr_invoice_no"),
            li.get("pr_supplier_bin"),
            li.get("pr_supplier_name"),
            str(li.get("pr_invoice_date") or "") or None,
            li.get("pr_taxable_amount_bdt"),
            li.get("pr_vat_amount_bdt"),
            li.get("sf_invoice_no"),
            str(li.get("sf_invoice_date") or "") or None,
            li.get("sf_taxable_amount_bdt"),
            li.get("sf_vat_amount_bdt"),
            li.get("match_status"),
            li.get("match_score"),
            li.get("ca_override"),
            li.get("ca_notes"),
        ])
        fill = _STATUS_FILLS.get(li.get("match_status", ""))
        if fill is not None:
            for col in range(1, len(_LINE_HEADERS) + 1):
                ws2.cell(row=ws2.max_row, column=col).fill = fill

    _autosize(ws2)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_reconciliation(
    *,
    reconciliation_id: UUID,
    tenant_id: UUID,
) -> bytes:
    """Load reconciliation + line items, return XLSX bytes.

    Raises ReconciliationNotFoundError if no row matches the id, and
    DocumentTenantMismatchError if the row belongs to a different tenant
    (defence in depth — RLS would also block this in normal flows).
    """
    supabase = get_supabase_admin()

    def _load() -> tuple[dict, list[dict]]:
        res = (
            supabase.table("vat_reconciliations")
            .select("*")
            .eq("id", str(reconciliation_id))
            .limit(1)
            .execute()
        )
        rows = res.data or []
        if not rows:
            raise ReconciliationNotFoundError(str(reconciliation_id))
        header = rows[0]
        if str(header["tenant_id"]) != str(tenant_id):
            raise DocumentTenantMismatchError(str(reconciliation_id))

        items = (
            supabase.table("recon_line_items")
            .select("*")
            .eq("reconciliation_id", str(reconciliation_id))
            .execute()
        )
        return header, items.data or []

    header, line_items = await asyncio.to_thread(_load)
    return build_xlsx_workbook(header=header, line_items=line_items)
