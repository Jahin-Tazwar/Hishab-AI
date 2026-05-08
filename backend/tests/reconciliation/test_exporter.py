"""Unit tests for the XLSX exporter (pure builder function)."""
from io import BytesIO

from openpyxl import load_workbook

from app.reconciliation.exporter import build_xlsx_workbook


def _sample_header():
    return {
        "id": "11111111-1111-1111-1111-111111111111",
        "period_start": "2024-01-01",
        "period_end": "2024-01-31",
        "status": "completed",
        "total_invoices": 3,
        "matched_exact": 1,
        "matched_fuzzy": 1,
        "partial_match": 0,
        "no_match": 1,
        "total_vat_claimed_bdt": "30000.00",
        "safe_itc_bdt": "20000.00",
        "at_risk_itc_bdt": "10000.00",
    }


def _sample_line_items():
    return [
        {
            "pr_invoice_no": "INV-001", "pr_supplier_bin": "123456789",
            "pr_supplier_name": "Acme", "pr_invoice_date": "2024-01-05",
            "pr_taxable_amount_bdt": "10000.00", "pr_vat_amount_bdt": "1500.00",
            "sf_invoice_no": "INV-001", "sf_invoice_date": "2024-01-05",
            "sf_taxable_amount_bdt": "10000.00", "sf_vat_amount_bdt": "1500.00",
            "match_status": "exact", "match_score": "1.00",
            "ca_override": None, "ca_notes": None,
        },
        {
            "pr_invoice_no": "INV-002", "pr_supplier_bin": "234567890",
            "pr_supplier_name": "Bravo", "pr_invoice_date": "2024-01-10",
            "pr_taxable_amount_bdt": "5000.00", "pr_vat_amount_bdt": "750.00",
            "sf_invoice_no": None, "sf_invoice_date": None,
            "sf_taxable_amount_bdt": None, "sf_vat_amount_bdt": None,
            "match_status": "no_match", "match_score": "0.00",
            "ca_override": None, "ca_notes": None,
        },
    ]


def test_build_xlsx_returns_bytes_with_two_sheets():
    data = build_xlsx_workbook(
        header=_sample_header(),
        line_items=_sample_line_items(),
    )
    assert isinstance(data, bytes) and len(data) > 0

    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Summary", "Line Items"]


def test_summary_sheet_includes_aggregate_numbers():
    data = build_xlsx_workbook(
        header=_sample_header(),
        line_items=_sample_line_items(),
    )
    wb = load_workbook(BytesIO(data))
    summary = wb["Summary"]
    # Build a simple {label: value} dict from column A/B
    rows = {summary.cell(r, 1).value: summary.cell(r, 2).value
            for r in range(1, summary.max_row + 1)}
    assert rows["Total Invoices"] == 3
    assert rows["Matched Exact"] == 1
    assert rows["Safe ITC (BDT)"] == "20000.00"


def test_line_items_sorted_no_match_first():
    data = build_xlsx_workbook(
        header=_sample_header(),
        line_items=_sample_line_items(),
    )
    wb = load_workbook(BytesIO(data))
    items = wb["Line Items"]
    # Header row at row 1; first data row should be the no_match invoice
    assert items.cell(2, 1).value == "INV-002"
    assert items.cell(2, 11).value == "no_match"
    assert items.cell(3, 1).value == "INV-001"
    assert items.cell(3, 11).value == "exact"


def test_empty_line_items_still_produces_workbook():
    data = build_xlsx_workbook(header=_sample_header(), line_items=[])
    wb = load_workbook(BytesIO(data))
    items = wb["Line Items"]
    # Headers only — no data rows
    assert items.max_row == 1
