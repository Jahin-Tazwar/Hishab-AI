"""Rendering tests. .docx is always tested; .pdf is gated on `soffice`."""
import io
import shutil
from decimal import Decimal

import pytest
from docx import Document

from app.notices.rendering import (
    PdfRendererUnavailableError, render_docx, render_pdf,
)


DRAFT = {
    "language": "bn",
    "body_html": "<p>প্রথম প্যারা</p><p>দ্বিতীয় প্যারা <sup class=\"citation\" data-ref=\"CIT-1\">[1]</sup></p>",
    "appendix_json": {
        "rows": [
            {"label": "Claimed ITC", "value_bdt": "50000.00", "note": None},
            {"label": "Allowed ITC", "value_bdt": "40000.00", "note": "per NBR"},
        ],
    },
    "citations": [
        {"corpus_chunk_id": "00000000-0000-0000-0000-0000000000aa",
         "source_ref": "Section 46", "snippet": "Conditions for ITC…",
         "paragraph_idx": 1},
    ],
}
NOTICE_META = {
    "notice_no": "NBR/VAT/1/2026", "notice_date": "2026-05-01",
    "taxpayer_bin": "001234567",
    "period_start": "2026-04-01", "period_end": "2026-04-30",
}
TENANT_META = {"firm_name": "Demo CA Co.", "address": "Dhaka, Bangladesh"}


def test_render_docx_contains_body_paragraphs_and_appendix():
    out = render_docx(draft=DRAFT, notice=NOTICE_META, tenant=TENANT_META)
    doc = Document(io.BytesIO(out))
    text_all = "\n".join(p.text for p in doc.paragraphs)
    assert "প্রথম প্যারা" in text_all
    assert "Demo CA Co." in text_all
    assert len(doc.tables) >= 1
    appendix_text = "\n".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells
    )
    assert "Claimed ITC" in appendix_text
    assert "50000.00" in appendix_text


def test_render_docx_includes_citations_section():
    out = render_docx(draft=DRAFT, notice=NOTICE_META, tenant=TENANT_META)
    doc = Document(io.BytesIO(out))
    text_all = "\n".join(p.text for p in doc.paragraphs)
    assert "Section 46" in text_all


@pytest.mark.skipif(
    shutil.which("soffice") is None,
    reason="LibreOffice (soffice) not on PATH; PDF rendering unavailable",
)
def test_render_pdf_emits_pdf_bytes():
    docx_bytes = render_docx(draft=DRAFT, notice=NOTICE_META, tenant=TENANT_META)
    pdf_bytes = render_pdf(docx_bytes)
    assert pdf_bytes[:4] == b"%PDF"


def test_render_pdf_raises_when_soffice_missing(monkeypatch):
    monkeypatch.setattr(
        "app.notices.rendering._soffice_path", lambda: None,
    )
    with pytest.raises(PdfRendererUnavailableError):
        render_pdf(b"not a real docx")
